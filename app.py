"""Web interface for the spreadsheet cleaner.

Run with:
    streamlit run app.py

Uploaded files are only ever held in memory; nothing is written to disk.

Set DEMO_MODE to "true" (environment variable or Streamlit secret) to cap uploads
for a public demo. Without it there are no limits.
"""

import csv
import io
import os
import zipfile
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook
from streamlit import config

from clean import clean, detect_headings, read_file, write_workbook
from make_demo_data import build_demo_files

HERE = Path(__file__).resolve().parent
# The first of these that exists is shown in the page header.
LOGO_FILES = ("assets/logo.svg", "assets/logo.png")

NO_KEY = "None (only remove exact duplicates)"
PREVIEW_ROWS = 100
XLSX_MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

DEMO_MAX_FILES = 3
DEMO_MAX_ROWS = 200
# Caps what a visitor can make the server hold and parse; the row limit alone is only checked after parsing.
DEMO_MAX_MB = 1
# An .xlsx is a zip; a tiny one can unpack to gigabytes, so the unpacked size is capped too.
DEMO_MAX_UNPACKED_MB = 10
# Byte limits don't bound parsing work: a few bytes can describe a sheet millions of cells wide
# or tall. These cap the shape itself, and are checked by streaming before anything is loaded.
DEMO_MAX_COLUMNS = 50
DEMO_MAX_SHEET_ROWS = 1000
DEMO_MESSAGE = "This is a demo. For a version tailored to your data, contact info.sendovia@gmail.com"


def demo_mode():
    """True when DEMO_MODE is "true" in the environment or in Streamlit secrets."""
    value = os.environ.get("DEMO_MODE")
    if value is None:
        try:
            value = st.secrets.get("DEMO_MODE")
        except Exception:
            # No secrets file, which is the normal case when running locally.
            value = None
    return str(value).strip().lower() == "true"


def as_sources(files):
    """Turn (name, bytes) pairs into the (name, open file) pairs clean.py expects."""
    return [(name, io.BytesIO(content)) for name, content in files]


@st.cache_data(show_spinner=False)
def demo_files():
    return sorted(build_demo_files().items())


@st.cache_data(show_spinner=False, max_entries=20)
def headings_for(files):
    return detect_headings(as_sources(files))


def enforce_demo_upload_cap():
    """Make the server itself refuse oversized uploads in demo mode.

    The uploader's own size limit is only checked in the visitor's browser. This sets the
    limit the server applies to every upload request, whatever sent it.
    """
    try:
        config.set_option("server.maxUploadSize", DEMO_MAX_MB)
    except Exception:
        # Still covered by the size check in over_demo_limits, which runs before any parsing.
        pass


def unpacked_size(content):
    """Total size an .xlsx would unpack to, read from the zip directory without unpacking it."""
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        return sum(entry.file_size for entry in archive.infolist())


def oversized_shape(name, content):
    """Describe how a file's rows/columns exceed the demo's shape limits, or return None.

    Streams the file and stops at the first row that breaks a limit, so the cost is bounded
    however large the sheet claims to be.
    """
    workbook = None
    try:
        if name.lower().endswith(".csv"):
            rows = csv.reader(io.StringIO(content.decode("latin-1"), newline=""))
        else:
            workbook = load_workbook(io.BytesIO(content), read_only=True)
            rows = workbook.worksheets[0].iter_rows(values_only=True)
        for count, row in enumerate(rows, start=1):
            if len(row) > DEMO_MAX_COLUMNS:
                return f"{name} has more than {DEMO_MAX_COLUMNS} columns, the most the demo allows."
            if count > DEMO_MAX_SHEET_ROWS:
                return f"{name} has far more than {DEMO_MAX_ROWS} rows, the most the demo allows per file."
    except Exception:
        return f"{name} could not be read."
    finally:
        if workbook is not None:
            workbook.close()
    return None


def over_demo_limits(uploads):
    """Describe how the uploads exceed the demo limits, or return None if they don't.

    The count and size checks come first so an oversized upload is never parsed.
    """
    if len(uploads) > DEMO_MAX_FILES:
        return f"You uploaded {len(uploads)} files; the demo allows up to {DEMO_MAX_FILES}."
    for upload in uploads:
        if upload.size > DEMO_MAX_MB * 1024 * 1024:
            return f"{upload.name} is larger than {DEMO_MAX_MB} MB, the most the demo allows per file."
    for upload in uploads:
        if not upload.name.lower().endswith(".csv"):
            try:
                too_big = unpacked_size(upload.getvalue()) > DEMO_MAX_UNPACKED_MB * 1024 * 1024
            except zipfile.BadZipFile:
                return f"{upload.name} is not a valid .xlsx file."
            if too_big:
                return f"{upload.name} unpacks to more than {DEMO_MAX_UNPACKED_MB} MB, the most the demo allows."
    for upload in uploads:
        problem = oversized_shape(upload.name, upload.getvalue())
        if problem:
            return problem
    for upload in uploads:
        try:
            rows = len(read_file(upload.name, io.BytesIO(upload.getvalue())))
        except Exception:
            continue  # unreadable files are reported when cleaning
        if rows > DEMO_MAX_ROWS:
            return f"{upload.name} has {rows} rows; the demo allows up to {DEMO_MAX_ROWS} rows per file."
    return None


def run_clean(files, key, dayfirst):
    table, summary, per_file, skipped = clean(as_sources(files), key=key, dayfirst=dayfirst)
    workbook = io.BytesIO()
    write_workbook(table, summary, per_file, workbook)
    return {"table": table, "summary": summary, "per_file": per_file, "skipped": skipped, "workbook": workbook.getvalue()}


def find_logo():
    for name in LOGO_FILES:
        if (HERE / name).is_file():
            return str(HERE / name)
    return None


def show_result(result):
    summary = dict(result["summary"])
    merged_label = next((label for label in summary if label.startswith("Rows merged by key")), None)

    st.download_button(
        "Download cleaned.xlsx",
        result["workbook"],
        file_name="cleaned.xlsx",
        mime=XLSX_MIME,
        type="primary",
        icon=":material/download:",
    )

    cards = [("Files read", "Files read"), ("Rows in", "Rows in"), ("Duplicates removed", "Duplicates removed")]
    if merged_label:
        cards.append(("Rows merged", merged_label))
    cards.append(("Rows out", "Rows out"))
    for column, (title, label) in zip(st.columns(len(cards)), cards):
        column.metric(title, summary[label], border=True)

    notes = [
        f"{summary[label]} {text}"
        for label, text in [
            ("Blank rows removed", "blank rows removed"),
            ("Dates not recognized (left as-is)", "dates not recognized (left as-is)"),
            ("Phone numbers not recognized (left as-is)", "phone numbers not recognized (left as-is)"),
        ]
        if summary[label]
    ]
    if notes:
        st.caption(" · ".join(notes))
    for name, error in result["skipped"]:
        st.warning(f"Skipped {name}: {error}")

    table = result["table"]
    shown = min(len(table), PREVIEW_ROWS)
    st.markdown("**Preview**")
    st.dataframe(table.head(PREVIEW_ROWS), hide_index=True, alt="Preview of the cleaned table")
    if len(table) > PREVIEW_ROWS:
        st.caption(f"Showing the first {shown} of {len(table)} rows. The download has all of them.")
    else:
        st.caption(f"All {len(table)} rows, {len(table.columns)} columns.")
    with st.expander("Rows read per file"):
        st.dataframe(
            pd.DataFrame(result["per_file"], columns=["File", "Rows read"]),
            hide_index=True,
            alt="Rows read from each file",
        )


st.set_page_config(page_title="Spreadsheet Cleaner", page_icon=":material/table_view:", layout="wide")
logo = find_logo()
if logo:
    st.logo(logo, size="large")
st.title("Spreadsheet Cleaner", anchor=False)
st.write(
    "Combine messy customer lists into one clean Excel file: matching column headings, "
    "tidy dates and phone numbers, duplicates removed."
)

DEMO = demo_mode()
if DEMO:
    enforce_demo_upload_cap()


def upload_step():
    """Step 1. Returns the files to clean as (name, bytes) pairs, or None if there are none yet."""
    uploads = st.file_uploader(
        "Upload .xlsx or .csv files",
        type=["xlsx", "csv"],
        accept_multiple_files=True,
        max_upload_size=DEMO_MAX_MB if DEMO else None,
    )
    if DEMO:
        st.caption(
            f"Demo limits: up to {DEMO_MAX_FILES} files; {DEMO_MAX_ROWS} rows, {DEMO_MAX_COLUMNS} columns "
            f"and {DEMO_MAX_MB} MB per file."
        )
    if st.button(
        "Load demo files",
        icon=":material/science:",
        help="Five made-up customer lists with messy headings, dates and duplicates.",
    ):
        st.session_state["demo"] = True

    if uploads:
        problem = over_demo_limits(uploads) if DEMO else None
        if problem:
            st.error(f"{problem} {DEMO_MESSAGE}")
            return None
        # Alphabetical, like the command-line tool, so "first value wins" behaves the same.
        return tuple(sorted((upload.name, upload.getvalue()) for upload in uploads))
    if st.session_state.get("demo"):
        files = tuple(demo_files())
        st.info("Using the demo files: " + ", ".join(name for name, _ in files))
        if st.button("Remove demo files"):
            st.session_state["demo"] = False
            st.rerun()
        return files
    return None


def options_step(files, headings):
    """Step 2. Returns the signature of the current files and options."""
    key_choice = st.selectbox(
        "Merge rows that share the same value in this column",
        [NO_KEY] + headings,
        help="For example, choose Email to get one row per customer. Blanks in one file are filled from another.",
    )
    dayfirst = st.checkbox("Dates are day-first (03/04/2024 means 3 April, not March 4)")
    key = None if key_choice == NO_KEY else key_choice

    # Identifies exactly what was cleaned, so an old result is never shown against new inputs.
    signature = (files, key, dayfirst)
    stored = st.session_state.get("result")
    up_to_date = bool(stored and stored[0] == signature)

    if st.button("Clean", type="secondary" if up_to_date else "primary", icon=":material/cleaning_services:"):
        try:
            with st.spinner("Cleaning..."):
                st.session_state["result"] = (signature, run_clean(files, key, dayfirst))
        except ValueError as error:
            st.session_state.pop("result", None)
            st.error(str(error))
        else:
            st.rerun()
    return signature


def page():
    with st.container(border=True):
        st.subheader(":material/upload_file: Step 1 · Upload", anchor=False)
        files = upload_step()
        headings = headings_for(files) if files else []
        if files and not headings:
            st.error("None of those files could be read. Check they are real .xlsx or .csv files.")
            files = None

    signature = None
    with st.container(border=True):
        st.subheader(":material/tune: Step 2 · Options", anchor=False)
        if files:
            signature = options_step(files, headings)
        else:
            st.caption("Add files in step 1 to choose your options.")

    with st.container(border=True):
        st.subheader(":material/download: Step 3 · Download", anchor=False)
        stored = st.session_state.get("result")
        if signature and stored and stored[0] == signature:
            show_result(stored[1])
        elif signature and stored:
            st.caption("Files or options have changed. Click Clean to update the result.")
        elif signature:
            st.caption("Click Clean in step 2 to see the summary, a preview, and your download.")
        else:
            st.caption("Your cleaned file will appear here.")


page()
if DEMO:
    st.divider()
    st.caption(DEMO_MESSAGE)
