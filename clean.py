#!/usr/bin/env python3
"""Merge every .xlsx and .csv file in a folder into one cleaned workbook.

Usage:
    python clean.py                 # reads ./input, writes ./output/cleaned.xlsx
    python clean.py --dayfirst      # read 03/04/2024 as 3 April instead of March 4
    python clean.py --key email     # merge rows that share the same email into one
"""

import argparse
import re
import sys
from collections import Counter
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent

PHONE_HINTS = ("phone", "mobile", "cell", "tel", "fax")
DATE_HINTS = ("date", "dob", "birthday")
EMAIL_HINTS = ("email",)

UNAMBIGUOUS_FORMATS = [
    "%Y-%m-%d", "%Y/%m/%d", "%Y.%m.%d",
    "%d-%b-%Y", "%d-%b-%y", "%d %b %Y", "%d %B %Y",
    "%b %d, %Y", "%B %d, %Y", "%b %d %Y", "%B %d %Y",
]
MONTH_FIRST_FORMATS = ["%m/%d/%Y", "%m-%d-%Y", "%m.%d.%Y", "%m/%d/%y", "%m-%d-%y"]
DAY_FIRST_FORMATS = ["%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y", "%d/%m/%y", "%d-%m-%y"]

TIME_SUFFIX = re.compile(r"[T\s]\d{1,2}:\d{2}(:\d{2})?(\.\d+)?(\s*[AaPp][Mm])?$")
ORDINAL = re.compile(r"(?<=\d)(st|nd|rd|th)\b", re.IGNORECASE)
PHONE_EXTENSION = re.compile(r"^(.*?)(?:\s*(?:ext\.?|x|#)\s*(\d+))?$", re.IGNORECASE)


def header_key(name):
    """'First Name', 'first_name' and ' FIRSTNAME ' all become 'firstname'."""
    return re.sub(r"[^a-z0-9]", "", str(name).lower())


def header_label(name):
    """Readable version of a header: 'signup_date' / 'SignupDate' -> 'Signup Date'."""
    text = re.sub(r"[_\-]+", " ", str(name))
    text = re.sub(r"(?<=[a-z0-9])(?=[A-Z])", " ", text)
    return re.sub(r"\s+", " ", text).strip().title()


def clean_text(value):
    return re.sub(r"\s+", " ", str(value)).strip()


def parse_date(value, dayfirst=False):
    """Return the date as YYYY-MM-DD, or None if it isn't a recognizable date."""
    text = TIME_SUFFIX.sub("", value)
    text = ORDINAL.sub("", text)
    text = re.sub(r"(?<=[A-Za-z])\.", "", text)
    text = re.sub(r"\s+", " ", text).strip()
    ambiguous = DAY_FIRST_FORMATS + MONTH_FIRST_FORMATS if dayfirst else MONTH_FIRST_FORMATS + DAY_FIRST_FORMATS
    for fmt in UNAMBIGUOUS_FORMATS + ambiguous:
        try:
            return datetime.strptime(text, fmt).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return None


def clean_phone(value):
    """Return a 10-digit number as (555) 123-4567, or None if it isn't one."""
    number, extension = PHONE_EXTENSION.match(value).groups()
    digits = re.sub(r"\D", "", re.sub(r"\.0$", "", number))
    if len(digits) == 11 and digits.startswith("1"):
        digits = digits[1:]
    if len(digits) != 10:
        return None
    formatted = f"({digits[:3]}) {digits[3:6]}-{digits[6:]}"
    return f"{formatted} x{extension}" if extension else formatted


def read_file(name, source):
    """Read one spreadsheet with every cell as text.

    source is a path or an open binary file; name decides whether it is read as CSV or Excel.
    """
    if name.lower().endswith(".csv"):
        for encoding in ("utf-8-sig", "cp1252", "latin-1"):
            if hasattr(source, "seek"):
                source.seek(0)
            try:
                df = pd.read_csv(source, dtype=str, keep_default_na=False, encoding=encoding)
                break
            except UnicodeDecodeError:
                continue
    else:
        df = pd.read_excel(source, dtype=str, keep_default_na=False)
    return df.fillna("").map(clean_text)


def normalize_columns(df, labels):
    """Rename columns to their matching key, merging any that collide within a file."""
    columns = {}
    for name in df.columns:
        series = df[name]
        if str(name).startswith("Unnamed:") and (series == "").all():
            continue
        key = header_key(name)
        if not key:
            continue
        labels.setdefault(key, Counter())[header_label(name)] += 1
        if key in columns:
            columns[key] = columns[key].where(columns[key] != "", series)
        else:
            columns[key] = series
    return pd.DataFrame(columns, index=df.index)


def column_label(labels, key):
    """The most common spelling of a heading; ties go to the one with more words."""
    return max(labels[key].items(), key=lambda item: (item[1], len(item[0].split())))[0]


def load_files(sources):
    """Read (name, path-or-file) pairs into frames with matched columns.

    Returns (frames, labels, per_file, skipped); unreadable files are skipped, not fatal.
    """
    labels = {}
    frames, per_file, skipped = [], [], []
    for name, source in sources:
        try:
            df = read_file(name, source)
        except Exception as error:
            skipped.append([name, str(error)])
            continue
        frames.append(normalize_columns(df, labels))
        per_file.append([name, len(df)])
    return frames, labels, per_file, skipped


def detect_headings(sources):
    """The column headings the cleaned table will have, in order."""
    frames, labels, _, _ = load_files(sources)
    if not frames:
        return []
    return [column_label(labels, key) for key in pd.concat(frames).columns]


def is_date_column(key, values, dayfirst):
    if any(hint in key for hint in DATE_HINTS):
        return True
    filled = [v for v in values if v]
    if not filled:
        return False
    parsed = sum(parse_date(v, dayfirst) is not None for v in filled)
    return parsed / len(filled) >= 0.8


def standardize(series, convert):
    """Apply convert to each filled cell; leave cells it can't handle untouched."""
    unrecognized = 0
    result = []
    for value in series:
        converted = convert(value) if value else value
        if converted is None:
            unrecognized += 1
            converted = value
        result.append(converted)
    return pd.Series(result, index=series.index), unrecognized


def merge_on_key(df, key):
    """Collapse rows sharing a key value into one, keeping the first non-empty value per column."""
    group = df[key].str.casefold()
    # Rows with no key value can't be matched to anything, so each stays on its own.
    group = group.where(group != "", pd.Series("\0" + df.index.astype(str), index=df.index))
    merged = df.replace("", pd.NA).groupby(group, sort=False).first()
    return merged.fillna("").reset_index(drop=True)


def write_workbook(df, summary, per_file, output_path):
    """Write the two-sheet workbook to a path or an in-memory buffer."""
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="305496")

    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        df.replace("", None).to_excel(writer, sheet_name="Cleaned Data", index=False)
        data_sheet = writer.sheets["Cleaned Data"]
        data_sheet.freeze_panes = "A2"
        data_sheet.auto_filter.ref = data_sheet.dimensions

        sheet = writer.book.create_sheet("Summary")
        sheet.append(["Metric", "Value"])
        for row in summary:
            sheet.append(row)
        sheet.append([])
        sheet.append(["File", "Rows read"])
        file_header_row = sheet.max_row
        for row in per_file:
            sheet.append(row)

        header_rows = {"Cleaned Data": [1], "Summary": [1, file_header_row]}
        for ws in (data_sheet, sheet):
            for row_number in header_rows[ws.title]:
                for cell in ws[row_number]:
                    cell.font = header_font
                    cell.fill = header_fill
            for index, column in enumerate(ws.columns, start=1):
                width = max((len(str(cell.value)) for cell in column if cell.value is not None), default=0)
                ws.column_dimensions[get_column_letter(index)].width = min(max(width + 2, 10), 50)


def clean(sources, key=None, dayfirst=False):
    """Merge and clean (name, path-or-file) pairs.

    Returns (table, summary, per_file, skipped). Raises ValueError if nothing could be
    read or the key column doesn't exist.
    """
    frames, labels, per_file, skipped = load_files(sources)
    if not frames:
        raise ValueError("None of the files could be read.")

    merged = pd.concat(frames, ignore_index=True).fillna("")
    rows_in = len(merged)

    blank = (merged == "").all(axis=1)
    merged = merged[~blank]

    merge_key = None
    if key:
        merge_key = header_key(key)
        if merge_key not in merged.columns:
            available = ", ".join(column_label(labels, column) for column in merged.columns)
            raise ValueError(f"No column called '{key}' to use as the key. Columns found: {available}")

    bad_dates = bad_phones = 0
    for column in merged.columns:
        if any(hint in column for hint in EMAIL_HINTS):
            merged[column] = merged[column].str.lower()
        elif any(hint in column for hint in PHONE_HINTS):
            merged[column], count = standardize(merged[column], clean_phone)
            bad_phones += count
        elif is_date_column(column, merged[column], dayfirst):
            merged[column], count = standardize(merged[column], lambda v: parse_date(v, dayfirst))
            bad_dates += count

    # Rows that differ only in upper/lower case count as duplicates; the first one is kept.
    duplicate = merged.apply(lambda column: column.str.casefold()).duplicated()
    merged = merged[~duplicate].reset_index(drop=True)

    key_merged = 0
    if merge_key:
        before = len(merged)
        merged = merge_on_key(merged, merge_key)
        key_merged = before - len(merged)

    merged.columns = [column_label(labels, column) for column in merged.columns]

    summary = [
        ["Files read", len(frames)],
        ["Rows in", rows_in],
        ["Blank rows removed", int(blank.sum())],
        ["Duplicates removed", int(duplicate.sum())],
        ["Rows out", len(merged)],
    ]
    if merge_key:
        summary.insert(4, [f"Rows merged by key ({key})", key_merged])
    summary += [
        ["Dates not recognized (left as-is)", bad_dates],
        ["Phone numbers not recognized (left as-is)", bad_phones],
    ]
    if skipped:
        summary.append(["Files skipped (could not be read)", ", ".join(name for name, _ in skipped)])
    return merged, summary, per_file, skipped


def main():
    parser = argparse.ArgumentParser(description="Merge and clean the spreadsheets in a folder.")
    parser.add_argument("--input", type=Path, default=HERE / "input", help="folder of .xlsx/.csv files (default: input)")
    parser.add_argument("--output", type=Path, default=HERE / "output" / "cleaned.xlsx", help="file to write (default: output/cleaned.xlsx)")
    parser.add_argument("--dayfirst", action="store_true", help="read 03/04/2024 as 3 April rather than March 4")
    parser.add_argument("--key", help="merge rows that share this column's value into one row (e.g. --key email)")
    args = parser.parse_args()

    if not args.input.is_dir():
        sys.exit(f"Could not find the input folder: {args.input}")
    paths = sorted(
        p for p in args.input.iterdir()
        if p.suffix.lower() in (".xlsx", ".csv") and not p.name.startswith(("~$", "."))
    )
    if not paths:
        sys.exit(f"No .xlsx or .csv files found in {args.input}")

    try:
        merged, summary, per_file, skipped = clean([(p.name, p) for p in paths], key=args.key, dayfirst=args.dayfirst)
    except ValueError as error:
        sys.exit(str(error))
    for name, error in skipped:
        print(f"  Skipped {name}: {error}")
    for name, rows in per_file:
        print(f"  Read {name}: {rows} rows")

    args.output.parent.mkdir(parents=True, exist_ok=True)
    try:
        write_workbook(merged, summary, per_file, args.output)
    except PermissionError:
        sys.exit(f"Could not save {args.output}. If it is open in Excel, close it and run again.")

    print()
    for label, value in summary:
        print(f"  {label}: {value}")
    print(f"\nSaved {args.output}")


if __name__ == "__main__":
    main()
